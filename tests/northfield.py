"""Fictional running example (spec Appendix A): Northfield University, Lot 1 student desks."""
import io

import docx
import pymupdf

SPEC_V1 = {
    1: [("1. Introduction", "Northfield University invites tenders for teaching-space furniture in two lots. "
         "Lot 1 covers student desks and Lot 2 covers seating. This specification forms part of the tender pack "
         "and must be read together with the Instructions to Tenderers.")],
    4: [("2.1 Quantities", "The supplier shall supply 200 student desks for Lot 1. Desks shall be delivered "
         "to the Main Campus teaching block.")],
    5: [("2.3 Product standard", "Desks shall comply with BS EN 1729-2 or an equivalent standard accepted by the "
         "University. Bidders claiming equivalence must provide test evidence from an accredited laboratory.")],
    6: [("3.2 Delivery", "The supplier shall deliver all desks within 30 days of order. Delivery shall be made "
         "between 08:00 and 16:00 on Working Days as defined in Appendix B.")],
    7: [("3.4 Installation", "Installation of all desks is included in the contract price. The supplier shall "
         "remove all packaging from site.")],
}


def pdf(pages: dict[int, list[tuple[str, str]]], last_page: int | None = None) -> bytes:
    doc = pymupdf.open()
    for pno in range(1, (last_page or max(pages)) + 1):
        page = doc.new_page()
        y = 72
        for heading, body in pages.get(pno, []):
            page.insert_text((72, y), heading, fontsize=13)
            y += 24
            box = pymupdf.Rect(72, y, 520, y + 300)
            page.insert_textbox(box, body, fontsize=10)
            y += 30 + 14 * (len(body) // 80 + 1)
        if pno not in pages:
            page.insert_text((72, 72), f"Page {pno} intentionally left blank for layout.", fontsize=10)
    return doc.tobytes()


def scanned_pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 400, 400), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    return doc.tobytes()


def itt_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Instructions to Tenderers", 0)
    d.add_heading("5.1 Experience", 1)
    d.add_paragraph("Bidders must provide details of at least one comparable contract for 150 or more items "
                    "delivered in the last three years, with a customer reference, using the form at Appendix C.")
    d.add_heading("6.2 Sustainability", 1)
    d.add_paragraph("The share of recycled material in the desks will be scored and is worth 10% of the quality score.")
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Criterion", "Weight"), ("Quality", "60%"), ("Price", "40%")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def release(rid="029001-2026", tag="tender", date="2026-10-01T09:00:00Z", deadline="2026-10-30T12:00:00Z",
            spec_url="https://portal.test/Spec_Desks_v1.pdf", itt_url="https://portal.test/login") -> dict:
    return {
        "ocid": "ocds-test-northfield", "id": rid, "date": date, "tag": [tag],
        "buyer": {"id": "GB-TEST-NORTHFIELD", "name": "Northfield University"},
        "tender": {
            "title": "Student desks and seating", "status": "active", "description": "Teaching-space furniture.",
            "legalBasis": {"id": "2023/54"},
            "tenderPeriod": {"endDate": deadline},
            "lots": [{"id": "1", "title": "Student desks", "description": "Supply and installation of student desks.",
                      "status": "active"},
                     {"id": "2", "title": "Seating", "status": "active"}],
            "items": [{"id": "1", "relatedLot": "1", "quantity": 200,
                       "additionalClassifications": [{"scheme": "CPV", "id": "39160000"}],
                       "deliveryAddresses": [{"region": "UKE"}]},
                      {"id": "2", "relatedLot": "2", "additionalClassifications": [{"scheme": "CPV", "id": "39110000"}],
                       "deliveryAddresses": [{"region": "UKE"}]}],
            "selectionCriteria": {"criteria": [{"type": "economic", "description": "Public liability insurance",
                                                "minimum": "At least £5m public liability insurance."}]},
            "documents": [
                {"id": "D-1", "documentType": "technicalSpecifications", "title": "Specification", "url": spec_url},
                {"id": "D-2", "documentType": "biddingDocuments", "title": "Instructions to Tenderers", "url": itt_url},
                {"id": "D-3", "documentType": "biddingDocuments", "title": "Pricing schedule", "url": itt_url},
                {"id": rid, "documentType": "tenderNotice", "noticeType": "UK4",
                 "url": f"https://www.find-tender.service.gov.uk/Notice/{rid}"},
            ],
        },
    }
